"""Categories (K-13) and simple data lists (K-06)."""

import csv
import io
from collections.abc import Sequence
from uuid import UUID

from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ConflictError, NotFoundError, PermissionDeniedError, ValidationFailedError
from app.modules.audit.service import AuditService
from app.modules.identity.models import User
from app.modules.knowledge.models import ItemKind, KnowledgeCategory, KnowledgeItem, KnowledgeListRow
from app.modules.knowledge.permissions import KnowledgePolicy, is_knowledge_admin
from app.modules.knowledge.schemas import CategoryIn, CategoryUpdate


class CategoryService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.audit = AuditService(session)

    @staticmethod
    def _require_admin(user: User) -> None:
        if not is_knowledge_admin(user):
            raise PermissionDeniedError("Only knowledge admins can change categories.")

    async def get(self, category_id: UUID) -> KnowledgeCategory:
        category = await self.session.get(KnowledgeCategory, category_id)
        if category is None:
            raise NotFoundError("We could not find that category.")
        return category

    async def create(self, data: CategoryIn, user: User) -> KnowledgeCategory:
        self._require_admin(user)
        if data.parent_id:
            await self.get(data.parent_id)
        position = int(
            await self.session.scalar(
                select(func.count()).where(
                    KnowledgeCategory.parent_id == data.parent_id if data.parent_id else KnowledgeCategory.parent_id.is_(None)
                )
            )
            or 0
        )
        category = KnowledgeCategory(**data.model_dump(), position=position)
        self.session.add(category)
        await self.session.flush()
        self.audit.record("knowledge.category_created", "knowledge_category", category.id, actor_id=user.id, name=data.name)
        return category

    async def update(self, category_id: UUID, data: CategoryUpdate, user: User) -> KnowledgeCategory:
        """Rename or move. Items and their history follow automatically."""
        self._require_admin(user)
        category = await self.get(category_id)
        if data.name:
            category.name = data.name
        if data.head_id is not None:
            category.head_id = data.head_id
        if data.move_to_root:
            category.parent_id = None
        elif data.parent_id:
            if data.parent_id == category.id or await self._is_descendant(data.parent_id, category.id):
                raise ValidationFailedError("A category cannot be moved inside itself.")
            category.parent_id = data.parent_id
        self.audit.record("knowledge.category_changed", "knowledge_category", category.id, actor_id=user.id)
        return category

    async def _is_descendant(self, candidate: UUID, ancestor: UUID) -> bool:
        current = await self.session.get(KnowledgeCategory, candidate)
        while current is not None and current.parent_id:
            if current.parent_id == ancestor:
                return True
            current = await self.session.get(KnowledgeCategory, current.parent_id)
        return False

    async def merge(self, category_id: UUID, into_id: UUID, user: User) -> int:
        self._require_admin(user)
        if category_id == into_id:
            raise ValidationFailedError("Choose a different category to merge into.")
        source, target = await self.get(category_id), await self.get(into_id)
        moved = await self.session.execute(
            update(KnowledgeItem).where(KnowledgeItem.category_id == source.id).values(category_id=target.id)
        )
        await self.session.execute(
            update(KnowledgeCategory).where(KnowledgeCategory.parent_id == source.id).values(parent_id=target.id)
        )
        await self.session.delete(source)
        self.audit.record("knowledge.category_merged", "knowledge_category", target.id, actor_id=user.id, merged=source.name)
        return int(moved.rowcount or 0)  # type: ignore[attr-defined]

    async def delete(self, category_id: UUID, user: User) -> None:
        self._require_admin(user)
        category = await self.get(category_id)
        used = await self.session.scalar(select(func.count()).where(KnowledgeItem.category_id == category.id))
        children = await self.session.scalar(select(func.count()).where(KnowledgeCategory.parent_id == category.id))
        if used or children:
            raise ConflictError("This category still has items. Merge it into another category instead.")
        await self.session.delete(category)


class ListService:
    """Any team can keep a register as a simple table: import, filter, row history (audit)."""

    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.policy = KnowledgePolicy(session)

    async def _list_item(self, item_id: UUID, user: User, *, edit: bool) -> KnowledgeItem:
        item, rights = await self.policy.get_readable(item_id, user)
        if item.kind != ItemKind.LIST.value:
            raise ValidationFailedError("This item is not a data list.")
        if edit and not rights.edit:
            raise PermissionDeniedError("Only the owner and editors can change this list.")
        return item

    async def rows(self, item_id: UUID, user: User, filters: dict[str, str]) -> tuple[list[str], Sequence[KnowledgeListRow]]:
        item = await self._list_item(item_id, user, edit=False)
        rows = (
            await self.session.scalars(
                select(KnowledgeListRow).where(KnowledgeListRow.item_id == item.id).order_by(KnowledgeListRow.position)
            )
        ).all()
        if filters:
            rows = [r for r in rows if all(str(r.data.get(k, "")).lower() == v.lower() for k, v in filters.items())]
        return item.list_columns, rows

    async def add_row(self, item_id: UUID, data: dict[str, object], user: User) -> KnowledgeListRow:
        item = await self._list_item(item_id, user, edit=True)
        clean = {c: data.get(c, "") for c in item.list_columns}
        position = int(await self.session.scalar(select(func.count()).where(KnowledgeListRow.item_id == item.id)) or 0)
        row = KnowledgeListRow(item_id=item.id, data=clean, position=position, updated_by_id=user.id)
        self.session.add(row)
        await self.session.flush()
        AuditService(self.session).record("knowledge.list_row_added", "knowledge_list_row", row.id, actor_id=user.id, data=clean)
        return row

    async def update_row(self, item_id: UUID, row_id: UUID, data: dict[str, object], user: User) -> KnowledgeListRow:
        item = await self._list_item(item_id, user, edit=True)
        row = await self.session.get(KnowledgeListRow, row_id)
        if row is None or row.item_id != item.id:
            raise NotFoundError("We could not find that row.")
        before = dict(row.data)
        row.data = {**row.data, **{k: v for k, v in data.items() if k in item.list_columns}}
        row.updated_by_id = user.id
        AuditService(self.session).record(
            "knowledge.list_row_changed", "knowledge_list_row", row.id, actor_id=user.id, before=before, after=row.data
        )  # row history
        return row

    async def delete_row(self, item_id: UUID, row_id: UUID, user: User) -> None:
        item = await self._list_item(item_id, user, edit=True)
        row = await self.session.get(KnowledgeListRow, row_id)
        if row is None or row.item_id != item.id:
            raise NotFoundError("We could not find that row.")
        AuditService(self.session).record(
            "knowledge.list_row_deleted", "knowledge_list_row", row.id, actor_id=user.id, data=row.data
        )
        await self.session.delete(row)

    async def import_file(self, item_id: UUID, filename: str, content: bytes, user: User) -> int:
        """Build the list from a spreadsheet: the first row becomes the columns."""
        item = await self._list_item(item_id, user, edit=True)
        if filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            table = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
        elif filename.lower().endswith(".csv"):
            table = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
        else:
            raise ValidationFailedError("Please upload a .xlsx or .csv file.")
        table = [r for r in table if any(c.strip() for c in r)]
        if len(table) < 2:
            raise ValidationFailedError("The file needs a header row and at least one row.")
        headers = [h.strip() or f"Column {i + 1}" for i, h in enumerate(table[0])]
        item.list_columns = headers
        existing = (await self.session.scalars(select(KnowledgeListRow).where(KnowledgeListRow.item_id == item.id))).all()
        for row in existing:
            await self.session.delete(row)
        for position, values in enumerate(table[1:]):
            self.session.add(
                KnowledgeListRow(
                    item_id=item.id,
                    position=position,
                    updated_by_id=user.id,
                    data={h: (values[i] if i < len(values) else "") for i, h in enumerate(headers)},
                )
            )
        AuditService(self.session).record(
            "knowledge.list_imported", "knowledge_item", item.id, actor_id=user.id, rows=len(table) - 1
        )
        return len(table) - 1
