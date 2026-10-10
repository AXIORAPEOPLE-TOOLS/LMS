import csv
import io
from collections.abc import Sequence
from uuid import UUID

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AuthenticationError, ConflictError, ValidationFailedError
from app.core.security import create_access_token, hash_password, verify_password
from app.modules.ai.provider import RoleCandidate
from app.modules.ai.service import get_assistant
from app.modules.audit.service import AuditService
from app.modules.identity.models import Department, SystemRole, User
from app.modules.identity.repository import DepartmentRepository, UserRepository
from app.modules.identity.schemas import (
    DepartmentIn,
    MeOut,
    PreferencesUpdate,
    SignupChoice,
    SignupIn,
    SignupOptionsOut,
    UploadCommitIn,
    UploadCommitOut,
    UploadPreviewOut,
    UploadRowOut,
    UserCreate,
    UserUpdate,
)
from app.modules.skills.models import JobRole
from app.modules.skills.repository import JobRoleRepository
from app.shared.events import DomainEvent, event_bus

VALID_ROLES = {r.value for r in SystemRole}


class AuthService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.users = UserRepository(session)
        self.audit = AuditService(session)

    async def login(self, email: str, password: str) -> str:
        user = await self.users.by_email(email)
        if user is None or not verify_password(password, user.password_hash) or not user.is_active:
            raise AuthenticationError("That email and password do not match.")
        self.audit.record("auth.login", "user", user.id, actor_id=user.id)
        return create_access_token(user.id)

    # ---- self sign-up: name, email, password, job role (the role brings its skills) ----------
    async def signup_options(self) -> SignupOptionsOut:
        from app.modules.settings.service import SettingsService

        settings = SettingsService(self.session)
        allowed = bool(await settings.get("identity.allow_self_signup"))
        domains = list(await settings.get("identity.signup_email_domains"))
        if not allowed:
            return SignupOptionsOut(allowed=False, email_domains=[], roles=[], departments=[])
        roles = await JobRoleRepository(self.session).list(order_by=JobRole.name)
        departments = await DepartmentRepository(self.session).list(order_by=Department.name)
        return SignupOptionsOut(
            allowed=True,
            email_domains=domains,
            roles=[SignupChoice(id=r.id, name=r.name, group=r.family) for r in roles],
            departments=[SignupChoice(id=d.id, name=d.name) for d in departments],
        )

    async def signup(self, data: SignupIn) -> str:
        options = await self.signup_options()
        if not options.allowed:
            raise ValidationFailedError("Sign-up is turned off. Ask HR to add you.")
        domain = data.email.rsplit("@", 1)[-1].lower()
        if options.email_domains and domain not in {d.lower() for d in options.email_domains}:
            raise ValidationFailedError(f"Use your work email ({', '.join(options.email_domains)}).")
        if await JobRoleRepository(self.session).get(data.job_role_id) is None:
            raise ValidationFailedError("Choose your job role from the list.")
        if data.department_id and await DepartmentRepository(self.session).get(data.department_id) is None:
            raise ValidationFailedError("Choose your department from the list.")
        user = await UserService(self.session).create_user(
            UserCreate(
                email=data.email,
                full_name=data.full_name.strip(),
                password=data.password,
                job_role_id=data.job_role_id,
                department_id=data.department_id,
                system_roles=["learner"],
            ),
            actor=None,
        )
        self.audit.record("user.signed_up", "user", user.id, actor_id=user.id, job_role_id=str(data.job_role_id))
        return create_access_token(user.id)


class UserService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.users = UserRepository(session)
        self.departments = DepartmentRepository(session)
        self.roles = JobRoleRepository(session)
        self.audit = AuditService(session)

    async def me(self, user: User) -> MeOut:
        dept = await self.departments.get(user.department_id) if user.department_id else None
        role = await self.roles.get(user.job_role_id) if user.job_role_id else None
        manager = await self.users.get(user.manager_id) if user.manager_id else None
        from app.modules.knowledge.permissions import can_open_hub  # late import: knowledge depends on identity

        return MeOut.model_validate(user).model_copy(
            update={
                "department_name": dept.name if dept else None,
                "job_role_name": role.name if role else None,
                "manager_name": manager.full_name if manager else None,
                "can_open_knowledge": await can_open_hub(self.session, user),
            }
        )

    async def update_preferences(self, user: User, data: PreferencesUpdate) -> User:
        for key, value in data.model_dump(exclude_unset=True).items():
            if key == "reminder_channels" and value is not None:
                value = {**(user.reminder_channels or {}), **value}
            if key == "phone":
                value = (value or "").strip() or None
            setattr(user, key, value)
        await self.session.flush()
        return user

    async def list_users(self, query: str | None = None, department_id: UUID | None = None) -> Sequence[User]:
        where = []
        if query:
            where.append(User.full_name.ilike(f"%{query}%") | User.email.ilike(f"%{query}%"))
        if department_id:
            where.append(User.department_id == department_id)
        return await self.users.list(*where, order_by=User.full_name)

    async def create_user(self, data: UserCreate, actor: User | None) -> User:
        if await self.users.by_email(data.email):
            raise ConflictError("Someone with this email already exists.")
        self._check_roles(data.system_roles)
        payload = data.model_dump(exclude={"password"})
        user = await self.users.add_and_flush(User(**payload, password_hash=hash_password(data.password)))
        self.audit.record("user.created", "user", user.id, actor_id=actor.id if actor else None, email=user.email)
        await event_bus.publish(self.session, DomainEvent("user.created", {"user_id": user.id}))
        return user

    async def update_user(self, user_id: UUID, data: UserUpdate, actor: User) -> User:
        user = await self.users.get_or_raise(user_id)
        changes = data.model_dump(exclude_unset=True)
        if "system_roles" in changes:
            self._check_roles(changes["system_roles"])
        old_role = user.job_role_id
        for key, value in changes.items():
            setattr(user, key, value)
        await self.session.flush()
        self.audit.record("user.updated", "user", user.id, actor_id=actor.id, changes=list(changes))
        if "job_role_id" in changes and changes["job_role_id"] != old_role:
            await event_bus.publish(self.session, DomainEvent("user.role_changed", {"user_id": user.id}))
        return user

    @staticmethod
    def _check_roles(roles: list[str]) -> None:
        unknown = set(roles) - VALID_ROLES
        if unknown:
            raise ValidationFailedError(f"Unknown access roles: {', '.join(sorted(unknown))}")

    # ---- departments -----------------------------------------------------
    async def list_departments(self) -> Sequence[Department]:
        return await self.departments.list(order_by=Department.name)

    async def create_department(self, data: DepartmentIn, actor: User) -> Department:
        if await self.departments.by_name(data.name):
            raise ConflictError("This department already exists.")
        dept = await self.departments.add_and_flush(Department(**data.model_dump()))
        self.audit.record("department.created", "department", dept.id, actor_id=actor.id)
        return dept

    # ---- employee upload (A-02) -----------------------------------------
    async def preview_upload(self, filename: str, content: bytes) -> UploadPreviewOut:
        records = _read_people_file(filename, content)
        roles = await self.roles.list(order_by=JobRole.name)
        exact = {r.name.lower(): r for r in roles}
        candidates = [RoleCandidate(r.id, r.name) for r in roles]
        assistant = get_assistant()
        rows: list[UploadRowOut] = []
        seen_emails: set[str] = set()
        for index, record in enumerate(records, start=2):  # row 1 is the header
            name, email = record.get("full_name", ""), record.get("email", "").lower()
            title, dept = record.get("job_title", ""), record.get("department") or None
            base = {"row": index, "full_name": name, "email": email, "job_title": title, "department": dept}
            if not name or "@" not in email:
                rows.append(UploadRowOut(**base, status="error", message="Name and a valid email are needed."))
                continue
            if email in seen_emails or await self.users.by_email(email):
                rows.append(UploadRowOut(**base, status="error", message="This email is already in use."))
                continue
            seen_emails.add(email)
            role = exact.get(title.lower())
            if role:
                rows.append(UploadRowOut(**base, status="ready", job_role_id=role.id))
                continue
            match = await assistant.match_role(title, candidates) if title else None
            rows.append(
                UploadRowOut(
                    **base,
                    status="needs_role",
                    suggested_role_id=match.role_id if match else None,
                    suggested_role_name=match.role_name if match else None,
                    message=None if match else "Choose a role for this person.",
                )
            )
        return UploadPreviewOut(
            rows=rows,
            ready=sum(r.status == "ready" for r in rows),
            needs_role=sum(r.status == "needs_role" for r in rows),
            errors=sum(r.status == "error" for r in rows),
        )

    async def commit_upload(self, data: UploadCommitIn, actor: User) -> UploadCommitOut:
        created = skipped = 0
        for row in data.rows:
            if await self.users.by_email(row.email):
                skipped += 1
                continue
            dept = None
            if row.department:
                dept = await self.departments.by_name(row.department) or await self.departments.add_and_flush(
                    Department(name=row.department)
                )
            await self.create_user(
                UserCreate(
                    email=row.email,
                    full_name=row.full_name,
                    password=_temporary_password(),
                    employee_code=row.employee_code,
                    job_role_id=row.job_role_id,
                    department_id=dept.id if dept else None,
                ),
                actor,
            )
            created += 1
        self.audit.record("user.bulk_upload", "user", None, actor_id=actor.id, created=created, skipped=skipped)
        return UploadCommitOut(created=created, skipped=skipped)


HEADER_ALIASES = {
    "name": "full_name",
    "full name": "full_name",
    "employee name": "full_name",
    "email": "email",
    "email address": "email",
    "work email": "email",
    "job title": "job_title",
    "title": "job_title",
    "designation": "job_title",
    "role": "job_title",
    "department": "department",
    "team": "department",
    "employee id": "employee_code",
    "employee code": "employee_code",
}

TEMPLATE_HEADERS = ["Full name", "Email", "Job title", "Department", "Employee ID"]


def _read_people_file(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True).active
        table = [[str(c) if c is not None else "" for c in row] for row in sheet.iter_rows(values_only=True)]
    elif filename.lower().endswith(".csv"):
        table = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    else:
        raise ValidationFailedError("Please upload the template as .xlsx or .csv.")
    if not table:
        raise ValidationFailedError("The file is empty.")
    headers = [HEADER_ALIASES.get(h.strip().lower(), h.strip().lower()) for h in table[0]]
    if "email" not in headers or "full_name" not in headers:
        raise ValidationFailedError("The file needs at least the columns 'Full name' and 'Email'.")
    return [
        {headers[i]: (cell or "").strip() for i, cell in enumerate(row) if i < len(headers)}
        for row in table[1:]
        if any((cell or "").strip() for cell in row)
    ]


def _temporary_password() -> str:
    import secrets

    return secrets.token_urlsafe(16)


def people_template_csv() -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TEMPLATE_HEADERS)
    writer.writerow(["Asha Verma", "asha@yourcompany.com", "Admissions Counsellor", "Admissions", "E001"])
    return buffer.getvalue()
