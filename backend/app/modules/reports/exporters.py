"""Export strategies: CSV, Excel, PDF. Every export carries who ran it, when and with which filters."""

import csv
import io
from collections.abc import Sequence
from typing import Any, Protocol

from app.modules.reports.base import Column


class Exporter(Protocol):
    media_type: str
    extension: str

    def export(self, title: str, columns: Sequence[Column], rows: list[dict[str, Any]], watermark: str) -> bytes: ...


class CsvExporter:
    media_type = "text/csv"
    extension = "csv"

    def export(self, title: str, columns: Sequence[Column], rows: list[dict[str, Any]], watermark: str) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([f"{title} — {watermark}"])
        writer.writerow([c.label for c in columns])
        for row in rows:
            writer.writerow(["" if row.get(c.key) is None else row.get(c.key) for c in columns])
        return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM so Excel opens Hindi names correctly


class ExcelExporter:
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def export(self, title: str, columns: Sequence[Column], rows: list[dict[str, Any]], watermark: str) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([watermark])
        ws["A2"].font = Font(italic=True, color="666666")
        ws.append([c.label for c in columns])
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5C4E")
        for row in rows:
            ws.append([row.get(c.key) for c in columns])
        for index, column in enumerate(columns, start=1):
            width = max([len(column.label), *(len(str(r.get(column.key) or "")) for r in rows[:200])]) + 2
            ws.column_dimensions[ws.cell(row=3, column=index).column_letter].width = min(width, 50)
        ws.freeze_panes = "A4"
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


class PdfExporter:
    media_type = "application/pdf"
    extension = "pdf"

    def export(self, title: str, columns: Sequence[Column], rows: list[dict[str, Any]], watermark: str) -> bytes:
        from fpdf import FPDF

        def latin(value: Any) -> str:
            return str("" if value is None else value).encode("latin-1", errors="replace").decode("latin-1")

        pdf = FPDF(orientation="L", unit="mm", format="A4")
        pdf.set_auto_page_break(True, margin=12)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 8, latin(title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "I", 8)
        pdf.set_text_color(110, 110, 110)
        pdf.cell(0, 6, latin(watermark), new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
        width = (pdf.w - 20) / max(1, len(columns))
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(31, 92, 78)
        pdf.set_text_color(255, 255, 255)
        for c in columns:
            pdf.cell(width, 7, latin(c.label)[:30], border=0, fill=True)
        pdf.ln()
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 8)
        for i, row in enumerate(rows):
            pdf.set_fill_color(244, 246, 245) if i % 2 else pdf.set_fill_color(255, 255, 255)
            for c in columns:
                pdf.cell(width, 6, latin(row.get(c.key))[:40], border=0, fill=True)
            pdf.ln()
        return bytes(pdf.output())


EXPORTERS: dict[str, Exporter] = {"csv": CsvExporter(), "xlsx": ExcelExporter(), "pdf": PdfExporter()}
